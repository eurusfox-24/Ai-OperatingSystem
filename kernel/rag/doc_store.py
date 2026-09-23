import os
import time
import logging
import uuid
from typing import List, Dict, Any, Optional
import numpy as np
from kernel.core.azure_client import azure_client
from kernel.db.local_manager import db_manager

logger = logging.getLogger("doc_store")

DOCS_STORAGE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data", "documents"))
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".pptx", ".txt", ".md", ".json"}
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024

class DocumentIngestionEngine:
    """Ingests PDF, Word (.docx), Excel (.xlsx/.csv), PowerPoint (.pptx), TXT, and Markdown files
    into local SQLite storage with sqlite-vec vector embeddings."""

    def __init__(self):
        os.makedirs(DOCS_STORAGE_DIR, exist_ok=True)
        self.vector_store: List[Dict[str, Any]] = []
        self.documents_metadata: Dict[str, Dict[str, Any]] = {}

    def parse_file(self, file_path: str) -> str:
        """Extracts clean text content from PDF, Word, Excel, PowerPoint, TXT, or MD files."""
        ext = os.path.splitext(file_path)[1].lower()
        text_content = ""

        try:
            if ext == ".pdf":
                import pypdf
                reader = pypdf.PdfReader(file_path)
                text_content = "\n".join([page.extract_text() or "" for page in reader.pages])

            elif ext == ".docx":
                import docx
                doc = docx.Document(file_path)
                text_content = "\n".join([p.text for p in doc.paragraphs if p.text])

            elif ext in [".xlsx", ".csv"]:
                import pandas as pd
                if ext == ".csv":
                    df = pd.read_csv(file_path)
                else:
                    df = pd.read_excel(file_path)
                text_content = f"Data Summary:\nColumns: {list(df.columns)}\nRows Count: {len(df)}\nSample Data:\n{df.head(20).to_string()}"

            elif ext == ".pptx":
                import pptx
                prs = pptx.Presentation(file_path)
                slides_text = []
                for idx, slide in enumerate(prs.slides):
                    slide_shapes_text = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            slide_shapes_text.append(shape.text)
                    slides_text.append(f"--- Slide {idx+1} ---\n" + "\n".join(slide_shapes_text))
                text_content = "\n\n".join(slides_text)

            elif ext in [".txt", ".md", ".json"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text_content = f.read()

            else:
                raise ValueError(f"Unsupported document type: {ext or 'no extension'}")

            return text_content.strip()

        except Exception as e:
            logger.error(f"Error parsing document {file_path}: {e}")
            raise e

    def generate_ai_summary(self, file_name: str, text: str) -> str:
        """Generates a concise executive summary of the document using Azure OpenAI."""
        prompt = f"Summarize the following document titled '{file_name}' in 2 clear executive sentences highlighting key topics, numbers, and targets:\n\n{text[:3000]}"
        try:
            res = azure_client.chat_completion(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.3,
                max_tokens=150
            )
            return res.get("content", "Executive summary unavailable.").strip()
        except Exception as e:
            logger.error(f"Error generating summary for {file_name}: {e}")
            return "Document ingested into local embedded SQLite vector store."

    def ingest_document(self, file_name: str, content_bytes: bytes) -> Dict[str, Any]:
        """Saves raw file locally, extracts text, generates embeddings, and stores chunks + vectors in local SQLite."""
        if not file_name or os.path.basename(file_name) != file_name or file_name in {".", ".."}:
            return {"status": "error", "message": "A safe document filename is required."}
        ext = os.path.splitext(file_name)[1].lower()
        if ext not in SUPPORTED_EXTENSIONS:
            return {"status": "error", "message": f"Unsupported document type: {ext or 'no extension'}"}
        if not content_bytes:
            return {"status": "error", "message": "The uploaded document is empty."}
        if len(content_bytes) > MAX_DOCUMENT_BYTES:
            return {"status": "error", "message": "Document exceeds the 25 MiB upload limit."}
        # Parse a temporary copy first so a failed re-upload cannot destroy the
        # last known-good raw file.
        file_path = os.path.join(DOCS_STORAGE_DIR, file_name)
        temporary_path = os.path.join(DOCS_STORAGE_DIR, f".upload-{uuid.uuid4().hex}{ext}")
        with open(temporary_path, "wb") as f:
            f.write(content_bytes)

        try:
            text = self.parse_file(temporary_path)
        except Exception as exc:
            try:
                os.remove(temporary_path)
            except OSError:
                pass
            return {"status": "error", "message": f"Document could not be parsed: {exc}"}
        if not text:
            try:
                os.remove(temporary_path)
            except OSError:
                pass
            return {"status": "error", "message": "Extracted text was empty."}

        # 2. Chunk text into ~500 token segments
        chunks = [text[i:i+1500] for i in range(0, len(text), 1200)]
        storage_path = f"data/documents/{file_name}"

        indexed_payload: List[Dict[str, Any]] = []
        for idx, chunk in enumerate(chunks):
            embedding: Optional[List[float]] = None
            try:
                embedding = azure_client.get_embedding(chunk)
            except Exception as exc:
                # Lexical retrieval remains available when embeddings are not
                # configured or the provider is temporarily unavailable.
                logger.warning("Embedding unavailable for %s chunk %s: %s", file_name, idx, exc)

            indexed_payload.append({"text": chunk, "embedding": embedding})

        summary = self.generate_ai_summary(file_name, text)
        if not db_manager.replace_document_index(
            file_name=file_name,
            extension=ext,
            size_bytes=len(content_bytes),
            text_length=len(text),
            storage_path=storage_path,
            ai_summary=summary,
            chunks=indexed_payload,
        ):
            try:
                os.remove(temporary_path)
            except OSError:
                pass
            return {"status": "error", "message": "Document index could not be stored atomically."}
        try:
            os.replace(temporary_path, file_path)
        except OSError as exc:
            logger.error("Document was indexed but its raw file could not be finalized: %s", exc)
            return {"status": "error", "message": "Document index was stored but raw file finalization failed."}

        indexed_chunks = len(indexed_payload)
        self.vector_store = [item for item in self.vector_store if item.get("doc_name") != file_name]
        self.vector_store.extend(
            {
                "doc_name": file_name,
                "chunk_index": idx,
                "text": item["text"],
                "embedding": item["embedding"],
            }
            for idx, item in enumerate(indexed_payload)
            if item["embedding"] is not None
        )
        meta = {
            "file_name": file_name,
            "extension": ext,
            "size_bytes": len(content_bytes),
            "text_length": len(text),
            "chunks_indexed": indexed_chunks,
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "ai_summary": summary,
            "extracted_text": text,
            "storage_path": storage_path,
        }
        self.documents_metadata[file_name] = meta

        return {
            "status": "success",
            "document": meta
        }

    def list_ingested_documents(self) -> List[Dict[str, Any]]:
        """Returns metadata for all ingested documents from local SQLite."""
        db_docs = db_manager.get_all_documents()
        if db_docs:
            return db_docs
        return list(self.documents_metadata.values())

    def get_document_details(self, file_name: str) -> Optional[Dict[str, Any]]:
        """Returns complete extracted text and chunk breakdown for document inspection."""
        return db_manager.get_document_details(file_name)

    def delete_document(self, file_name: str) -> bool:
        """Delete database state, raw storage, and process-local caches together."""
        if not db_manager.delete_document(file_name):
            return False
        self.documents_metadata.pop(file_name, None)
        self.vector_store = [item for item in self.vector_store if item.get("doc_name") != file_name]
        file_path = os.path.abspath(os.path.join(DOCS_STORAGE_DIR, file_name))
        if os.path.commonpath([DOCS_STORAGE_DIR, file_path]) != DOCS_STORAGE_DIR:
            logger.error("Refusing to remove document outside storage: %s", file_path)
            return False
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
        except OSError as exc:
            logger.error("Database document was deleted but raw file removal failed for %s: %s", file_name, exc)
            return False
        return True

    def repair_missing_indexes(self) -> Dict[str, int]:
        """Rebuild lexical chunks for existing rows that falsely claim an index."""
        repaired = 0
        failed = 0
        for document in db_manager.get_all_documents():
            file_name = str(document.get("file_name") or "")
            details = db_manager.get_document_details(file_name)
            if not file_name or (details and details.get("chunks")):
                continue
            file_path = os.path.abspath(os.path.join(DOCS_STORAGE_DIR, file_name))
            if os.path.commonpath([DOCS_STORAGE_DIR, file_path]) != DOCS_STORAGE_DIR or not os.path.isfile(file_path):
                failed += 1
                continue
            try:
                text = self.parse_file(file_path)
                chunks = [text[i:i+1500] for i in range(0, len(text), 1200)]
                db_manager.clear_document_chunks(file_name)
                stored = sum(
                    db_manager.insert_chunk(file_name, index, chunk, None) >= 0
                    for index, chunk in enumerate(chunks)
                )
                if stored != len(chunks):
                    raise RuntimeError("Not all chunks were stored")
                db_manager.save_document_metadata(
                    file_name=file_name,
                    extension=str(document.get("extension") or os.path.splitext(file_name)[1].lower()),
                    size_bytes=int(document.get("size_bytes") or os.path.getsize(file_path)),
                    text_length=len(text),
                    chunks_indexed=stored,
                    storage_path=str(document.get("storage_path") or f"data/documents/{file_name}"),
                    ai_summary=str(document.get("ai_summary") or ""),
                )
                repaired += 1
            except Exception as exc:
                logger.error("Could not repair document index for %s: %s", file_name, exc)
                failed += 1
        return {"repaired": repaired, "failed": failed}

    def generate_notebook_briefing(self, notebook_id: str) -> str:
        """NotebookLM feature: Generates a comprehensive Briefing Document (Source Guide, FAQ, Timeline)."""
        nb = db_manager.get_notebook(notebook_id)
        if not nb or not nb.get("doc_names"):
            return "Notebook empty or not found."
            
        doc_summaries = []
        for doc_name in nb["doc_names"]:
            doc_meta = self.get_document_details(doc_name)
            if doc_meta:
                summary = doc_meta.get("ai_summary", "No summary.")
                doc_summaries.append(f"Document: {doc_name}\nSummary: {summary}")
                
        if not doc_summaries:
            return "No valid documents to brief."
            
        context_str = "\n\n".join(doc_summaries)
        
        prompt = (
            "You are a NotebookLM Source Guide generator.\n"
            "Based on the following document summaries, generate a comprehensive Briefing Document.\n"
            "Include:\n1. Executive Summary\n2. Key Entities\n3. FAQ (3-5 common questions)\n4. Brief Timeline or Action Plan\n\n"
            f"Sources:\n{context_str}"
        )
        
        try:
            res = azure_client.chat_completion(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.4,
                max_tokens=800
            )
            briefing_md = res.get("content", "Failed to generate briefing.")
            db_manager.save_notebook_briefing(notebook_id, briefing_md)
            return briefing_md
        except Exception as e:
            logger.error(f"Error generating briefing for notebook {notebook_id}: {e}")
            return f"Error generating briefing: {e}"

    def search_relevant_docs(
        self,
        query: str,
        top_k: int = 3,
        notebook_id: Optional[str] = None,
        document_names: Optional[List[str]] = None,
    ) -> str:
        """Searches only the requested project source scope when one is supplied."""
        if not self.vector_store and db_manager.get_db_stats().get("vec_chunks_count", 0) == 0:
            lexical_results = db_manager.search_document_chunks_lexically(
                query,
                top_k=top_k,
                notebook_id=notebook_id,
                document_names=document_names,
            )
            if not lexical_results:
                return "No internal documents uploaded yet."
            return "\n\n---\n\n".join(
                f"📄 [Source: {item['doc_name']} | Keyword match]\n{item['text_content']}"
                for item in lexical_results
            )

        try:
            query_embedding = azure_client.get_embedding(query)
            
            # Try local SQLite sqlite-vec KNN search first
            vec_results = db_manager.search_vectors(
                query_embedding,
                top_k=top_k,
                notebook_id=notebook_id,
                document_names=document_names,
            )
            if vec_results:
                formatted = []
                for item in vec_results:
                    distance = item.get("distance", 0.0)
                    doc_name = item.get("doc_name", "Unknown Document")
                    text = item.get("text_content", "")
                    # NotebookLM citation formatting
                    formatted.append(f"📄 [Source: {doc_name} | Distance: {distance:.4f}]\n{text}")
                return "\n\n---\n\n".join(formatted)

            # Fallback cosine distance calculation on in-memory store
            if not self.vector_store:
                lexical_results = db_manager.search_document_chunks_lexically(
                    query,
                    top_k=top_k,
                    notebook_id=notebook_id,
                    document_names=document_names,
                )
                if lexical_results:
                    return "\n\n---\n\n".join(
                        f"📄 [Source: {item['doc_name']} | Keyword match]\n{item['text_content']}"
                        for item in lexical_results
                    )
                return "No documents available for search."

            q_vec = np.array(query_embedding)
            scored_docs = []
            nb_docs = None
            if notebook_id:
                nb = db_manager.get_notebook(notebook_id)
                if nb:
                    nb_docs = nb.get("doc_names", [])
            if document_names is not None:
                requested_docs = {str(name) for name in document_names if str(name).strip()}
                candidate_docs = nb_docs if nb_docs is not None else requested_docs
                nb_docs = [name for name in candidate_docs if name in requested_docs]
                    
            for item in self.vector_store:
                if nb_docs is not None and item["doc_name"] not in nb_docs:
                    continue
                doc_vec = np.array(item["embedding"])
                norm_q = float(np.linalg.norm(q_vec))
                norm_doc = float(np.linalg.norm(doc_vec))
                if norm_q > 0 and norm_doc > 0:
                    similarity = float(np.dot(q_vec, doc_vec) / (norm_q * norm_doc))
                else:
                    similarity = 0.0
                scored_docs.append((similarity, item["doc_name"], item["text"]))

            scored_docs.sort(key=lambda x: x[0], reverse=True)
            top_results = scored_docs[:top_k]

            formatted = []
            for score, doc_name, text in top_results:
                formatted.append(f"📄 [Source: {doc_name} | Match Score: {score:.2f}]\n{text}")

            return "\n\n---\n\n".join(formatted)

        except Exception as e:
            logger.error(f"Document search error: {e}")
            lexical_results = db_manager.search_document_chunks_lexically(
                query,
                top_k=top_k,
                notebook_id=notebook_id,
                document_names=document_names,
            )
            if lexical_results:
                return "\n\n---\n\n".join(
                    f"📄 [Source: {item['doc_name']} | Keyword match]\n{item['text_content']}"
                    for item in lexical_results
                )
            return "Error querying local embedded vector document store."

doc_engine = DocumentIngestionEngine()
